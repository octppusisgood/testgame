# -*- coding: utf-8 -*-
"""通用 md → docx 转换（A4 横版）。用法: py -3 tests/md2docx.py <源.md> <目标docx路径> [标题]"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1]
DST = sys.argv[2]
TITLE = sys.argv[3] if len(sys.argv) > 3 else "游戏数值配置"

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

def set_cn_font(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = color

def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run(TITLE)
set_cn_font(r, 26, True, RGBColor(0x1F, 0x3B, 0x63))
doc.add_paragraph()

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip() or line.strip() == "---":
        i += 1
        continue
    if line.startswith("# "):
        i += 1
        continue
    if line.startswith("### "):
        h = doc.add_heading(level=2)
        r = h.add_run(line[4:].strip())
        set_cn_font(r, 13, True, RGBColor(0x2E, 0x5A, 0x88))
        i += 1
        continue
    if line.startswith("## "):
        h = doc.add_heading(level=1)
        r = h.add_run(line[3:].strip())
        set_cn_font(r, 16, True, RGBColor(0x1F, 0x3B, 0x63))
        i += 1
        continue
    if line.startswith("|"):
        tbl_lines = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            tbl_lines.append(lines[i].strip())
            i += 1
        rows = []
        for tl in tbl_lines:
            cells = [c.strip() for c in tl.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            rows.append(cells)
        if not rows:
            continue
        ncol = max(len(r_) for r_ in rows)
        table = doc.add_table(rows=len(rows), cols=ncol)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        fsize = 9.5 if ncol <= 6 else 8.5
        for ri, row in enumerate(rows):
            for ci in range(ncol):
                cell = table.cell(ri, ci)
                text = row[ci] if ci < len(row) else ""
                text = text.replace("~~", "")
                p = cell.paragraphs[0]
                r = p.add_run(text)
                if ri == 0:
                    set_cn_font(r, fsize, True, RGBColor(0xFF, 0xFF, 0xFF))
                    shade_cell(cell, "1F3B63")
                else:
                    set_cn_font(r, fsize)
                    if ri % 2 == 0:
                        shade_cell(cell, "EEF2F8")
        doc.add_paragraph()
        continue
    m = re.match(r"^(\d+)\.\s+(.*)", line)
    p = doc.add_paragraph()
    text = line
    if m:
        text = m.group(1) + ". " + m.group(2)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_cn_font(r, 10.5, True)
        else:
            r = p.add_run(part.replace("`", "").replace("~~", ""))
            set_cn_font(r, 10.5)
    i += 1

doc.save(DST)
print("saved:", DST)
